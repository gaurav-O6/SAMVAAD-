from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "SAMVAAD_research_paper_final.docx"
TEMPLATE = Path(r"C:\Users\Gaurav\Downloads\ws-vjcs.docx")

TITLE = "SAMVAAD: A Web-Based Multimodal Assistive Communication System for Sign, Speech, and Braille Interaction"
AUTHORS = [
    "Gaurav Thakare",
    "Priyanka Narawade",
    "Shreya Choure",
    "Krushnal Mahajan",
]
AFFILIATION = "SAMVAAD Project Team, Final Year Engineering Project"
CORRESPONDENCE = "Corresponding author: Gaurav Thakare"
HISTORY = "Received 4 May 2026\nRevised 4 May 2026"
KEYWORDS = (
    "assistive technology; multimodal interaction; sign language recognition; "
    "Braille OCR; accessibility; avatar-based communication."
)
CLASSIFICATION = "AMS Subject Classification: 68T07, 68U10, 68T45"

TABLE_ONE_ROWS = [
    ("Gesture sample labels", "41", "JSONL files under dataset/gesture_samples"),
    ("Recorded gesture samples", "1177", "Labeled landmark examples used for sample matching"),
    ("Braille image files", "43", "PNG image files under dataset"),
    ("Regression tests", "9 passed", "Verified by unittest discovery in the project virtual environment"),
    ("Backend stack", "Flask + CORS", "Implemented across app.py, templates/app.py, and requirements.txt"),
    ("Core vision stack", "MediaPipe / OpenCV / NumPy", "Implemented in sign_recog.py and samvaad_braille.py"),
]

TABLE_TWO_ROWS = [
    ("Braille character-image benchmark", "36 graded", "100.0%", "All labeled A-Z and 0-9 dataset checks passed in the bundled test set."),
    ("Braille auxiliary examples", "7 info", "Decoded", "Name and sentence images decoded successfully but were excluded from the graded accuracy figure."),
    ("Rule-only gesture classification", "1177 samples", "75.19%", "Geometric rules remain useful but exhibit class confusion on stored landmark samples."),
    ("Recorded-sample matcher", "1177 samples", "99.92%", "Shows a very strong fit on stored labeled samples, but not yet signer-independent generalization."),
    ("Combined gesture output", "1177 samples", "99.92%", "Hybrid rule plus sample matching aligns with nearly all stored samples and should be re-evaluated on held-out data."),
    ("Software regression tests", "9 tests", "Pass", "Route handling and Braille helper regression tests completed successfully."),
]

SECTIONS = [
    (
        "1. Introduction",
        [
            "Communication barriers remain a persistent challenge for people who rely on sign language, tactile reading systems, speech assistance, or combinations of these modalities. Many academic systems still concentrate on a single modality, which fragments the user experience and forces people to switch tools depending on context.",
            "SAMVAAD, expanded here as Sign and Multimodal Voice Assistive Access and Dialogue, was developed to investigate whether a lightweight web platform can unify sign-to-text, text-to-sign, speech-to-text, text-to-speech, and Braille image interpretation within one accessible interface. The system emphasizes deployability with common hardware such as a browser camera, standard image uploads, and commodity computing resources.",
            "The present paper documents the implemented prototype, its architecture, its current evidence base, and the technical choices that make the platform extensible for future academic evaluation and assistive-technology refinement.",
        ],
    ),
    (
        "2. Problem Statement and Motivation",
        [
            "Assistive communication workflows are often disconnected. A signer may need gesture recognition, a low-vision user may need Braille interpretation, and a classroom or household setting may require speech synthesis and text display at the same time. When these pathways are isolated, the interaction becomes slower, less inclusive, and more dependent on informal human mediation.",
            "SAMVAAD addresses this gap by treating sign recognition, Braille interpretation, speech services, and avatar-based playback as complementary accessibility channels within a single browser-driven environment. The underlying motivation is not merely to recognize isolated inputs, but to reduce transition cost between modes of communication.",
        ],
    ),
    (
        "3. Contributions",
        [
            "The main contribution of SAMVAAD lies in multimodal integration rather than in proposing a single stand-alone recognition algorithm. The implementation makes the following concrete contributions:",
        ],
    ),
    (
        "4. Related Work and Positioning",
        [
            "Accessibility research commonly progresses along separate tracks such as sign-language recognition, Braille document interpretation, speech assistance, and educational visualization. Sign-language studies often focus on sensor gloves, image classification, skeletal landmarks, or deep sequence models, whereas Braille studies emphasize embossed-dot localization, OCR, and assistive reading support.",
            "SAMVAAD is positioned at the intersection of these threads. Its research value is system-level integration: browser-based landmark capture, rule-guided sign interpretation, speech interaction, avatar playback, and image-based Braille recognition are combined in one deployable web application. This framing makes the project suitable for a systems-oriented accessibility paper while remaining transparent about the prototype stage of the work.",
        ],
    ),
    (
        "5. System Overview",
        [
            "SAMVAAD follows a browser-client and Flask-backend architecture. The browser handles camera access, MediaPipe-based hand tracking, text entry, speech services, and 3D avatar rendering. The backend serves the interface assets and exposes JSON endpoints for sign processing, gesture-sample management, common-gesture discovery, and Braille recognition.",
            "The sign-recognition pipeline begins with client-side hand landmarks. These landmarks are mirrored and normalized on the server before being passed to rule-based gesture classifiers for right-hand alphabetic and command gestures and left-hand numeric and response gestures. The backend also compares incoming landmarks against recorded example embeddings stored in JSONL files, after which a temporal stabilization layer confirms outputs only after repeated consistent observations.",
            "The Braille pipeline accepts an uploaded image, decodes it into an OpenCV array, improves contrast using CLAHE, suppresses noise through blurring and morphology, detects candidate dots from contours, segments dots into Braille cells, decodes the resulting patterns, and reconstructs user-facing text with optional spacing recovery.",
        ],
    ),
    (
        "6. Sign Recognition Module",
        [
            "The sign-recognition logic is implemented primarily in sign_recog.py and templates/app.py. The classifier is intentionally geometry-driven rather than model-heavy: it uses relative landmark distances, finger-state logic, orientation checks, and ambiguity thresholds to differentiate alphabetic and control gestures. This design keeps inference lightweight, explainable, and straightforward to debug.",
            "To reduce flicker and accidental outputs, the system employs a StableGesture mechanism. Candidate labels must persist across multiple frames and satisfy a majority condition within recent history before they are committed as outputs. This hold-to-confirm behavior is important for real-time browser interaction because even intentional static gestures exhibit frame-to-frame variation.",
            "The backend also supports recorded-sample matching. When labeled gesture samples are present in dataset/gesture_samples, SAMVAAD creates normalized embeddings from landmark coordinates and compares new inputs against nearby stored examples. This does not replace the rule-based recognizer; rather, it supplements it when the stored sample offers a stronger match.",
        ],
    ),
    (
        "7. Braille Recognition Module",
        [
            "The Braille module is implemented in samvaad_braille.py as a multi-stage image-processing pipeline. It accepts both file paths and image arrays, which makes it usable as a standalone command-line utility and as a backend component for uploaded browser images.",
            "Preprocessing is designed for practical imaging conditions such as low contrast, uneven illumination, compression artifacts, and inversion ambiguity. The pipeline converts images to grayscale, applies CLAHE for local contrast enhancement, smooths the result with Gaussian blur, uses adaptive thresholding to produce a binary image, performs morphological closing to repair small gaps, and auto-inverts the image when required.",
            "Dot detection relies on contour analysis with dynamic thresholds derived from the image itself. Candidate blobs are filtered by area, circularity, aspect ratio, and solidity before being interpreted as Braille dots. Cell segmentation then estimates inter-dot spacing, clusters rows, pairs columns, infers missing middle rows when necessary, and decodes letters, digits, punctuation, capitalization, and number indicators before optional post-processing restores spacing.",
        ],
    ),
    (
        "8. Multimodal User Interface and Avatar Playback",
        [
            "Beyond recognition, SAMVAAD includes interface features that make the platform communicative in both input and output directions. The browser uses Web Speech APIs for speech-to-text and text-to-speech, allowing users to move between spoken and written forms without leaving the application.",
            "For visual output, the system includes a 3D avatar player implemented with Three.js and animation assets stored under templates/animations. Text tokens and common commands can be mapped to avatar animations, supporting text-to-sign demonstrations and educational use. A learn mode and gesture recorder further strengthen the prototype as both an assistive and instructional environment.",
        ],
    ),
    (
        "9. Implementation Evidence and Current Prototype Status",
        [
            "The paper intentionally avoids unsupported deployment claims and instead summarizes evidence that can be verified directly from the repository. The current codebase already demonstrates broad multimodal coverage, but it should still be understood as a prototype pending larger benchmark studies and user-centered evaluation.",
        ],
    ),
    (
        "10. Preliminary Evaluation and Results",
        [
            "Preliminary validation was conducted using artifacts already present in the repository. The Braille module was evaluated on the bundled single-character image set, while the gesture-recognition pipeline was checked against the stored JSONL landmark samples used by the sample-matching component. The software regression suite was also executed in the project virtual environment.",
            "These results should be interpreted carefully. The Braille benchmark provides a useful first correctness check on the included labeled image set. By contrast, the gesture-sample evaluation is an internal validation on stored examples and therefore does not establish signer-independent generalization. It remains useful, however, because it quantifies the difference between the raw rule-based classifier and the hybrid sample-assisted pathway implemented in the current codebase.",
        ],
    ),
    (
        "11. Discussion",
        [
            "The engineering direction of SAMVAAD favors interpretability, portability, and low-cost deployment over large opaque models. That choice is appropriate for an academic prototype because it allows the recognition logic to be explained, debugged, and extended without specialized hardware.",
            "At the same time, the current architecture highlights the main technical challenges ahead. Rule-based gesture recognition remains sensitive to viewpoint, signer variation, and lighting conditions. Braille decoding quality still depends on image quality and dataset diversity. The avatar pathway demonstrates expressive output, but richer sentence-level translation would require more advanced sequencing and language handling than the present token-level design.",
            "Taken together, these properties make SAMVAAD strongest as a multimodal assistive-systems prototype with explainable modules and deployment-aware design choices. That framing is more credible than presenting the system as a fully benchmarked production recognizer.",
        ],
    ),
    (
        "12. Limitations and Future Work",
        [
            "Several limitations remain before formal publication in a fully benchmark-driven venue. The repository does not yet provide a standardized held-out test set with accuracy, precision, recall, latency, or signer-variation statistics across all modalities. In addition, the project artifacts do not yet encode a user study, so claims about usability and accessibility benefit should remain appropriately cautious.",
            "Future work should therefore prioritize curated multimodal datasets, held-out evaluation protocols, broader sign vocabulary coverage, additional Braille benchmarks, sentence-level translation, and usability studies with representative participants. These steps would convert the present implementation-backed prototype into a stronger comparative research paper.",
        ],
    ),
    (
        "13. Conclusion",
        [
            "SAMVAAD demonstrates that a single web-based system can integrate sign recognition, speech interaction, Braille interpretation, and avatar-based visual output within one assistive communication platform. Even at the prototype stage, the implementation provides a concrete and extensible foundation for future work in multimodal accessibility, human-computer interaction, and deployment-oriented assistive technology research.",
        ],
    ),
]

BULLETS = [
    "A unified Flask-based accessibility interface that serves sign recognition, Braille recognition, learn mode, and 3D avatar playback from one application.",
    "A hybrid sign-recognition workflow that combines browser-side hand landmark extraction with server-side rule-based classification, sample matching, and temporal stabilization.",
    "A Braille recognition pipeline that performs image preprocessing, contour-based dot filtering, cell segmentation, decoding, and spacing recovery.",
    "An extensible local dataset design in which recorded gesture samples are stored as JSONL files and can be reused for prototype refinement without retraining a heavy model.",
    "A practical project structure with automated regression tests, making the system easier to demonstrate, maintain, and extend in an academic setting.",
]

REFERENCES = [
    "1. Bradski, G. The OpenCV Library. Dr. Dobb's Journal of Software Tools (2000).",
    "2. Lugaresi, C., Tang, J., Nash, H., et al. MediaPipe: A Framework for Building Perception Pipelines. arXiv:1906.08172 (2019).",
    "3. W3C. Web Speech API Specification and implementation guidance. Available at: https://www.w3.org/TR/webspeechapi/.",
    "4. Pallets Project. Flask Documentation. Available at: https://flask.palletsprojects.com/.",
    "5. Three.js Authors. Three.js Documentation. Available at: https://threejs.org/.",
    "6. World Health Organization. Global report on assistive technology (2022).",
]


def set_run_font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_body(document: Document) -> None:
    body = document._element.body
    section_props = body.sectPr
    for child in list(body):
        body.remove(child)
    if section_props is not None:
        body.append(section_props)


def style_lookup(document: Document, preferred: str, fallback: str = "Normal") -> str:
    try:
        document.styles[preferred]
        return preferred
    except KeyError:
        return fallback


def add_paragraph(document: Document, text: str = "", style: str = "Normal", align: WD_ALIGN_PARAGRAPH | None = None):
    paragraph = document.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def set_cell_margins(cell, top: int = 70, start: int = 110, bottom: int = 70, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{key}"))
        if element is None:
            element = OxmlElement(f"w:{key}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def apply_table_geometry(table, widths_inches: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_table_cell(cell, text: str, *, bold: bool = False, align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=8, bold=bold)


def build_title_block(document: Document) -> None:
    title_style = style_lookup(document, "Article Title")
    author_style = style_lookup(document, "Author")
    aff_style = style_lookup(document, "Affiliation")
    history_style = style_lookup(document, "History", fallback=aff_style)

    p = document.add_paragraph(style=title_style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.clear()
    run = p.add_run(TITLE)
    set_run_font(run)

    for index, author in enumerate(AUTHORS):
        p = document.add_paragraph(style=author_style)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.clear()
        run = p.add_run(author)
        set_run_font(run)

        p = document.add_paragraph(style=aff_style)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.clear()
        affiliation_text = AFFILIATION if index < len(AUTHORS) - 1 else f"{AFFILIATION}\n{CORRESPONDENCE}"
        run = p.add_run(affiliation_text)
        set_run_font(run)

    p = document.add_paragraph(style=history_style)
    p.clear()
    run = p.add_run(HISTORY)
    set_run_font(run)


def build_abstract(document: Document) -> None:
    abstract_style = style_lookup(document, "Abstract")
    keywords_style = style_lookup(document, "Keywords")

    p = document.add_paragraph(style=abstract_style)
    p.clear()
    body = p.add_run(
        "SAMVAAD is a multimodal assistive communication prototype designed to reduce interaction barriers across hearing, speech, and visual impairments. "
        "The system combines browser-based hand landmark capture, server-side sign classification, speech input and output, text-to-sign avatar playback, "
        "and image-based Braille recognition within a single web application. Its sign-recognition module integrates MediaPipe landmarks, rule-based gesture "
        "decoding, recorded-sample matching, and temporal stabilization, while the Braille module uses contrast enhancement, adaptive thresholding, contour-based "
        "dot detection, cell segmentation, and pattern decoding to convert embossed-dot images into readable text. Implemented through a Flask backend with "
        "lightweight regression tests and extensible local datasets, SAMVAAD contributes a practical accessibility platform that unifies multiple communication "
        "pathways inside one lightweight architecture. The present paper reports the system design, current implementation evidence, preliminary evaluation results, "
        "and the main limitations that should guide future benchmark-driven research."
    )
    set_run_font(body)

    p = document.add_paragraph(style=keywords_style)
    p.clear()
    lead = p.add_run("Keywords: ")
    set_run_font(lead, bold=True)
    body = p.add_run(KEYWORDS)
    set_run_font(body)

    p = document.add_paragraph(style=keywords_style)
    p.clear()
    run = p.add_run(CLASSIFICATION)
    set_run_font(run)


def add_section(document: Document, title: str, paragraphs: list[str]) -> None:
    heading_style = style_lookup(document, "Heading 1")
    text_style = style_lookup(document, "Text")
    indent_style = style_lookup(document, "Text Indent", fallback=text_style)

    heading = document.add_paragraph(style=heading_style)
    heading.clear()
    run = heading.add_run(title)
    set_run_font(run)

    for index, text in enumerate(paragraphs):
        paragraph = document.add_paragraph(style=text_style if index == 0 else indent_style)
        paragraph.clear()
        run = paragraph.add_run(text)
        set_run_font(run)


def add_bullets(document: Document, items: list[str]) -> None:
    bullet_style = style_lookup(document, "Bullet List", fallback="List Bullet")
    for item in items:
        paragraph = document.add_paragraph(style=bullet_style)
        paragraph.clear()
        run = paragraph.add_run(item)
        set_run_font(run)


def add_table_caption(document: Document, text: str) -> None:
    caption_style = style_lookup(document, "Table Caption", fallback="Normal")
    paragraph = document.add_paragraph(style=caption_style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run)


def add_implementation_table(document: Document) -> None:
    add_table_caption(document, "Table 1. Snapshot of implementation-backed evidence available in the current repository.")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    apply_table_geometry(table, [1.95, 1.25, 2.95])
    header = table.rows[0]
    set_repeat_table_header(header)
    headers = ["Artifact", "Observed Value", "Evidence in Repository"]
    for idx, text in enumerate(headers):
        set_table_cell(header.cells[idx], text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for artifact, value, evidence in TABLE_ONE_ROWS:
        row = table.add_row().cells
        set_table_cell(row[0], artifact)
        set_table_cell(row[1], value, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_table_cell(row[2], evidence)


def add_evaluation_table(document: Document) -> None:
    add_table_caption(document, "Table 2. Preliminary validation results obtained from the current repository artifacts.")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    apply_table_geometry(table, [1.85, 0.95, 0.85, 2.5])
    header = table.rows[0]
    set_repeat_table_header(header)
    headers = ["Module / Check", "Samples", "Result", "Interpretation"]
    for idx, text in enumerate(headers):
        set_table_cell(header.cells[idx], text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for check, samples, result, interpretation in TABLE_TWO_ROWS:
        row = table.add_row().cells
        set_table_cell(row[0], check)
        set_table_cell(row[1], samples, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_table_cell(row[2], result, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_table_cell(row[3], interpretation)


def add_references(document: Document) -> None:
    heading_style = style_lookup(document, "Heading 1")
    reference_style = style_lookup(document, "Reference", fallback="Normal")

    heading = document.add_paragraph(style=heading_style)
    heading.clear()
    run = heading.add_run("References")
    set_run_font(run)

    for reference in REFERENCES:
        paragraph = document.add_paragraph(style=reference_style)
        paragraph.clear()
        run = paragraph.add_run(reference)
        set_run_font(run)


def build_document() -> Document:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    document = Document(TEMPLATE)
    clear_body(document)
    build_title_block(document)
    build_abstract(document)

    for title, paragraphs in SECTIONS:
        add_section(document, title, paragraphs)
        if title == "3. Contributions":
            add_bullets(document, BULLETS)
        elif title == "9. Implementation Evidence and Current Prototype Status":
            add_implementation_table(document)
            paragraph = document.add_paragraph(style=style_lookup(document, "Text Indent", fallback="Text"))
            paragraph.clear()
            run = paragraph.add_run(
                "The current codebase supports right-hand alphabetic and command gestures such as A to Z, HELLO, STOP, and SPACE, alongside left-hand numeric and response gestures including 1 to 10, YES, and NO. "
                "The presence of recorded samples for 41 labels indicates that data collection is already underway for practical refinement. The passing unit test suite further provides a baseline level of confidence for route behavior and core Braille helpers."
            )
            set_run_font(run)
        elif title == "10. Preliminary Evaluation and Results":
            add_evaluation_table(document)
            paragraph = document.add_paragraph(style=style_lookup(document, "Text Indent", fallback="Text"))
            paragraph.clear()
            run = paragraph.add_run(
                "The preliminary results are encouraging. The Braille recognizer achieved 100.0% accuracy on the 36 graded single-character images present in the bundled dataset, while also successfully decoding several ungraded name and sentence images. "
                "For gesture recognition, the rule-only classifier matched 75.19% of the stored gesture samples, whereas the recorded-sample matcher and the combined hybrid output each matched 99.92% of the same stored sample set. "
                "These findings suggest that the sample-assisted pathway materially improves fit to the collected examples, while also confirming the need for held-out evaluation before broader generalization claims are made."
            )
            set_run_font(run)

    add_references(document)
    return document


def main() -> None:
    document = build_document()
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
